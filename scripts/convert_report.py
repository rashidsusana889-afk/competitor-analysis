#!/usr/bin/env python3
"""
报告格式转换脚本
将 Markdown 报告转换为 PDF 和 Word 格式
"""

import os
import sys
from pathlib import Path


def convert_md_to_pdf(md_path: str, output_path: str) -> bool:
    """将 Markdown 转换为 PDF"""
    try:
        import weasyprint
        from markdown import markdown
        
        # 读取 Markdown 文件
        with open(md_path, 'r', encoding='utf-8') as f:
            md_content = f.read()
        
        # 转换为 HTML
        html_content = markdown(md_content, extensions=['extra', 'codehilite'])
        
        # 添加样式
        styled_html = f"""
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                body {{ font-family: 'Helvetica Neue', Arial, sans-serif; margin: 40px; line-height: 1.6; }}
                h1 {{ color: #2c3e50; border-bottom: 2px solid #3498db; padding-bottom: 10px; }}
                h2 {{ color: #34495e; margin-top: 30px; }}
                h3 {{ color: #7f8c8d; }}
                code {{ background: #f5f5f5; padding: 2px 6px; border-radius: 3px; }}
                pre {{ background: #f5f5f5; padding: 15px; border-radius: 5px; overflow-x: auto; }}
                table {{ border-collapse: collapse; width: 100%; margin: 20px 0; }}
                th, td {{ border: 1px solid #ddd; padding: 12px; text-align: left; }}
                th {{ background: #3498db; color: white; }}
                tr:nth-child(even) {{ background: #f9f9f9; }}
            </style>
        </head>
        <body>
        {html_content}
        </body>
        </html>
        """
        
        # 生成 PDF
        weasyprint.HTML(styled_html).write_pdf(output_path)
        print(f"✅ PDF 生成成功: {output_path}")
        return True
        
    except ImportError as e:
        print(f"⚠️  缺少依赖: {e}")
        print("   安装命令: pip install weasyprint markdown")
        return False
    except Exception as e:
        print(f"❌ PDF 转换失败: {e}")
        return False


def convert_md_to_word(md_path: str, output_path: str) -> bool:
    """将 Markdown 转换为 Word"""
    try:
        from markdown import markdown
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        # 读取 Markdown 文件
        with open(md_path, 'r', encoding='utf-8') as f:
            md_content = f.read()
        
        # 转换为 HTML 再解析
        html = markdown(md_content, extensions=['extra'])
        
        # 创建 Word 文档
        doc = Document()
        
        # 添加标题
        title = doc.add_heading('运动类App竞品更新分析报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 简单解析 HTML
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html, 'html.parser')
        
        for element in soup.find_all(['h1', 'h2', 'h3', 'h4', 'p', 'ul', 'ol', 'li', 'table']):
            if element.name == 'h1':
                doc.add_heading(element.get_text(), level=1)
            elif element.name == 'h2':
                doc.add_heading(element.get_text(), level=2)
            elif element.name == 'h3':
                doc.add_heading(element.get_text(), level=3)
            elif element.name == 'p':
                doc.add_paragraph(element.get_text())
            elif element.name in ['ul', 'ol']:
                for li in element.find_all('li'):
                    doc.add_paragraph(li.get_text(), style='List Bullet')
            elif element.name == 'table':
                table = doc.add_table(rows=1, cols=len(element.find('tr').find_all('th')) if element.find('tr').find_all('th') else 2)
                table.style = 'Table Grid'
                
                # 表头
                header_row = table.rows[0]
                headers = element.find('tr').find_all('th') if element.find('tr').find_all('th') else element.find('tr').find_all('td')
                for i, header in enumerate(headers):
                    header_row.cells[i].text = header.get_text()
                
                # 数据行
                for row in element.find_all('tr')[1:]:
                    cells = row.find_all('td')
                    if cells:
                        row_cells = table.add_row()
                        for i, cell in enumerate(cells):
                            row_cells.cells[i].text = cell.get_text()
        
        # 保存 Word 文档
        doc.save(output_path)
        print(f"✅ Word 生成成功: {output_path}")
        return True
        
    except ImportError as e:
        print(f"⚠️  缺少依赖: {e}")
        print("   安装命令: pip install python-docx markdown beautifulsoup4")
        return False
    except Exception as e:
        print(f"❌ Word 转换失败: {e}")
        return False


def convert_all_reports(reports_dir: str = './reports'):
    """转换所有报告"""
    reports_path = Path(reports_dir)
    
    if not reports_path.exists():
        print(f"❌ 目录不存在: {reports_dir}")
        return
    
    # 创建输出目录
    pdf_dir = reports_path / 'pdf'
    word_dir = reports_path / 'word'
    pdf_dir.mkdir(exist_ok=True)
    word_dir.mkdir(exist_ok=True)
    
    # 查找所有 md 文件
    md_files = list(reports_path.glob('*.md'))
    
    if not md_files:
        print("❌ 未找到 Markdown 报告文件")
        return
    
    print(f"找到 {len(md_files)} 个报告文件")
    
    for md_file in md_files:
        print(f"\n处理: {md_file.name}")
        
        # PDF
        pdf_path = pdf_dir / f"{md_file.stem}.pdf"
        convert_md_to_pdf(str(md_file), str(pdf_path))
        
        # Word
        word_path = word_dir / f"{md_file.stem}.docx"
        convert_md_to_word(str(md_file), str(word_path))
    
    print("\n✅ 全部转换完成!")


if __name__ == "__main__":
    reports_dir = sys.argv[1] if len(sys.argv) > 1 else './reports'
    convert_all_reports(reports_dir)
